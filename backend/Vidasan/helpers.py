from django.conf import settings
from datetime import datetime
import openpyxl
import re
import os
import pandas as pd
from io import BytesIO
from django.http import HttpResponse

def createWorkOrder(siparis, user):
    data = {
        'definition': siparis.definition,
        'description': siparis.description,
        'amount': siparis.amount,
        'currentDate': datetime.now().strftime('%d.%m.%Y'),
        'isOEM': 'Evet' if siparis.isOEM else 'Hayır',
        'material': siparis.material,
        'materialNumber': siparis.materialNumber,
        'lotNo': siparis.orderNumber,
        'company': siparis.company,
    }

    for proc in [ 'press', 'byck', 'ovalama', 'sementasyon', 'kaplama', 'ambalaj' ]:
        machineNumber = getattr( siparis.activity, f'{proc}Machine' )
        machineNumber = machineNumber.name if machineNumber else ''
        data[ f'{proc}MachineNumber' ] = machineNumber

        amount = getattr( siparis.activity, f'{proc}Amount' ) or ''
        data[ f'{proc}Amount' ] = amount

        outputKg = getattr( siparis.activity, f'{proc}OutputKg' ) or ''
        data[ f'{proc}OutputKg' ] = outputKg

        wastageKg = getattr( siparis.activity, f'{proc}WastageKg' ) or ''
        data[ f'{proc}WastageKg' ] = wastageKg

        startDateTime = getattr( siparis.activity, f'{proc}StartDateTime' )
        startDate = startDateTime.strftime('%d.%m.%Y') if startDateTime else ''
        startTime = startDateTime.strftime('%H:%M') if startDateTime else ''
        data[ f'{proc}StartDate' ] = startDate
        data[ f'{proc}StartTime' ] = startTime

        finishDateTime = getattr( siparis.activity, f'{proc}FinishDateTime' )
        finishDate = finishDateTime.strftime('%d.%m.%Y') if finishDateTime else ''
        finishTime = finishDateTime.strftime('%H:%M') if finishDateTime else ''
        data[ f'{proc}FinishDate' ] = finishDate
        data[ f'{proc}FinishTime' ] = finishTime

    formPath = os.path.join(settings.BASE_DIR, 'static', 'Vidasan', 'isEmriFormu.xlsx')
    workbook = openpyxl.load_workbook(formPath)
    sheet = workbook.active

    # Define a regex pattern to find placeholders
    pattern = re.compile(r'\{\{(\w+)\}\}')

    # Iterate over the cells and replace placeholders
    for row in sheet.iter_rows():
        for cell in row:
            if cell.value and isinstance(cell.value, str):
                # Find all placeholders in the cell value
                matches = pattern.findall(cell.value)
                for match in matches:
                    # Get the replacement value from the dictionary or use empty string if not found
                    replacement = str( data.get(match, '') )
                    # Replace the placeholder with the corresponding value
                    cell.value = pattern.sub(lambda m: replacement if m.group(1) == match else m.group(0), cell.value)
    
    # Save the modified Excel to a BytesIO object
    excel_buffer = BytesIO()
    workbook.save(excel_buffer)
    excel_buffer.seek(0)

    return excel_buffer
